import sys, glob, os
from datetime import date
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAME, ROLL, CLASS = "Shravani Chavan", "23102A0055", "CMPN A"
PROJECT = "CI/CD Pipeline for a Parcel Tracking Portal"

# ---------- helpers ----------
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()

def add_screenshots(doc, task_no):
    folder = f"docs/screenshots/task-{task_no:02d}"
    files = sorted(glob.glob(folder + "/*.png") + glob.glob(folder + "/*.jpg"))
    if not files:
        doc.add_paragraph(f"(No screenshots found in {folder}. Add images there and re-run.)")
        return
    for i, f in enumerate(files, 1):
        doc.add_picture(f, width=Inches(6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(f"Figure {i}: {os.path.basename(f)}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

def build(task_no, title, branch, blocks, filename):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

    h = doc.add_heading(f"Task {task_no}: {title}", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_table(doc, ["Field", "Details"], [
        ["Project", PROJECT], ["Name", NAME], ["Roll No", ROLL], ["Class", CLASS],
        ["Branch", branch], ["Date", date.today().strftime("%d %B %Y")],
    ])

    for b in blocks:
        kind = b[0]
        if kind == "h":
            doc.add_heading(b[1], 1)
        elif kind == "h2":
            doc.add_heading(b[1], 2)
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "bullets":
            for x in b[1]: doc.add_paragraph(x, style="List Bullet")
        elif kind == "numbered":
            for x in b[1]: doc.add_paragraph(x, style="List Number")
        elif kind == "check":
            for x in b[1]: doc.add_paragraph("\u2611 " + x)
        elif kind == "table":
            add_table(doc, b[1], b[2])
        elif kind == "code":
            p = doc.add_paragraph(); r = p.add_run(b[1])
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif kind == "screenshots":
            add_screenshots(doc, task_no)

    os.makedirs("docs/reports", exist_ok=True)
    out = f"docs/reports/{filename}.docx"
    doc.save(out)
    print("Created:", out)

# ---------- TASK CONTENT ----------
TASKS = {}

TASKS[1] = dict(
    title="Problem Definition and Scope",
    branch="task/01-problem-definition",
    filename="task-01-problem-definition",
    blocks=[
        ("h", "1. Objective"),
        ("p", "To study the real-time need for a Parcel Tracking Portal, identify its users, pain points, stakeholders, constraints and measurable success criteria, and freeze a small 15-task MVP scope that will be delivered through a complete CI/CD pipeline."),
        ("h", "2. Problem Statement"),
        ("p", "Customers and courier staff often have no single, reliable place to see the live status of a parcel. Customers repeatedly contact support to ask where their parcel is, while staff update statuses manually in scattered sheets or messages. Delayed, stuck or failed deliveries are noticed late, which leads to poor customer experience and extra support workload."),
        ("p", "The Parcel Tracking Portal solves this by giving one web application where parcel events are recorded, searched and monitored, with delayed or exceptional parcels flagged automatically."),
        ("h", "3. Target Users"),
        ("table", ["User", "Need"], [
            ["Customer", "Check the current status and history of a parcel quickly"],
            ["Delivery agent", "Update parcel events (picked up, in transit, delivered, failed)"],
            ["Warehouse staff", "Record arrival and dispatch of parcels"],
            ["Admin / support team", "Monitor all parcels, see delays and handle exceptions"],
        ]),
        ("h2", "3.1 Existing Pain Points"),
        ("bullets", [
            "No real-time visibility of parcel status.",
            "Status updates are manual and inconsistent.",
            "No automatic alert for delayed or failed deliveries.",
            "Parcel data is scattered across sheets, messages and calls.",
            "High number of repeated status queries to support.",
            "No quick summary of delivered, in-transit and delayed parcels.",
        ]),
        ("h", "4. Stakeholders"),
        ("bullets", [
            "Customers (end users of tracking)",
            "Courier company management",
            "Warehouse and delivery teams",
            "Customer support team",
            "Developers / DevOps team (project team)",
            "Course instructor / evaluator",
        ]),
        ("h", "5. Constraints"),
        ("bullets", [
            "Academic timeline: 15 tasks to be completed in sequence.",
            "Only free and open-source tools (Git, GitHub, Maven, Jenkins, Docker, Selenium, Ansible).",
            "Development and testing on a single laptop environment.",
            "No integration with a real courier API or live GPS; data is entered manually or simulated.",
            "Small MVP: the focus is on the DevOps pipeline, not on a large application.",
        ]),
        ("h", "6. Measurable Success Criteria"),
        ("table", ["#", "Criterion", "How it is measured"], [
            ["1", "Parcel search returns a result in under 2 seconds", "Manual timing / Selenium test"],
            ["2", "Full event history of a parcel visible on one screen", "Status drill-down page check"],
            ["3", "Delayed or failed parcels flagged on the dashboard automatically", "Alert/exception view check"],
            ["4", "Build, test and deploy run automatically on every commit", "Jenkins pipeline run"],
            ["5", "At least 80% of Selenium tests pass on each build", "Jenkins test report"],
            ["6", "Application runs identically in a Docker container", "Container run and health check"],
            ["7", "Environment can be rebuilt using an Ansible/Puppet script", "Re-run shows idempotent result"],
        ]),
        ("h", "7. Frozen MVP Scope"),
        ("h2", "In scope"),
        ("bullets", [
            "Parcel data / event entry",
            "Searchable dashboard",
            "Summary indicators (total, in transit, delivered, delayed)",
            "Status drill-down (event history per parcel)",
            "Alert / exception view for delayed or failed parcels",
        ]),
        ("h2", "Out of scope"),
        ("bullets", [
            "Real GPS or live map tracking",
            "Payments and billing",
            "Mobile application",
            "Real courier API integration",
            "Email/SMS notifications",
        ]),
        ("h", "8. Technology Stack (planned)"),
        ("table", ["Area", "Tool"], [
            ["Application", "Java Servlets/JSP"],
            ["Build", "Maven"],
            ["Version control", "Git and GitHub"],
            ["CI / CD", "Jenkins (Jenkinsfile)"],
            ["Server", "Apache Tomcat"],
            ["Testing", "Selenium WebDriver with JUnit"],
            ["Containers", "Docker"],
            ["Configuration management", "Ansible"],
        ]),
        ("h", "9. 15-Task Plan"),
        ("table", ["Task", "Title", "Branch"], [
            ["1", "Problem Definition and Scope", "task/01-problem-definition"],
            ["2", "Agile Planning and DevOps Workflow", "task/02-agile-planning"],
            ["3", "Requirements, Architecture and Technology Setup", "task/03-architecture-setup"],
            ["4", "Git and GitHub Repository Initialization", "task/04-repo-init"],
            ["5", "Feature Development with Branching", "task/05-feature-branching"],
            ["6", "MVP Completion and Git Collaboration", "task/06-mvp-completion"],
            ["7", "Jenkins Installation and Continuous Integration Job", "task/07-jenkins-ci"],
            ["8", "Pipeline as Code and Server Deployment", "task/08-pipeline-as-code"],
            ["9", "Selenium Test Design and Local Execution", "task/09-selenium-tests"],
            ["10", "Continuous Testing in Jenkins", "task/10-continuous-testing"],
            ["11", "Docker Image and Container Lifecycle", "task/11-docker"],
            ["12", "Jenkins-Docker Continuous Deployment", "task/12-jenkins-docker-cd"],
            ["13", "Configuration Management Script", "task/13-config-management"],
            ["14", "Automated Provisioning and Reliability Validation", "task/14-provisioning"],
            ["15", "Final End-to-End Release, Documentation and Viva", "task/15-final-release"],
        ]),
        ("h", "10. Steps Performed"),
        ("numbered", [
            "Studied the real-time tracking need and listed users and pain points.",
            "Identified stakeholders and constraints.",
            "Defined measurable success criteria.",
            "Froze the MVP scope and the 15-task plan.",
            "Created branch task/01-problem-definition from develop.",
            "Wrote the report, committed it and raised a pull request into develop.",
        ]),
        ("h", "11. Git Commands Used"),
        ("code", "git checkout develop\ngit checkout -b task/01-problem-definition\ngit add docs\ngit commit -m \"docs: add Task 1 problem definition and scope report\"\ngit push -u origin task/01-problem-definition"),
        ("h", "12. Screenshots"),
        ("screenshots",),
        ("h", "13. Deliverables Checklist"),
        ("check", [
            "Problem statement",
            "Target users and pain points",
            "Stakeholders and constraints",
            "Measurable success criteria",
            "Frozen 15-task MVP scope",
        ]),
        ("h", "14. Issues Faced and Fixes"),
        ("bullets", [
            "Linux-style mkdir -p failed in Windows PowerShell; used mkdir docs\\reports instead.",
            "Clone command was entered with a placeholder username; corrected with the real GitHub username.",
        ]),
        ("h", "15. Conclusion"),
        ("p", "The problem, users, scope and success criteria of the Parcel Tracking Portal are now clearly defined and frozen. This gives a stable base for Agile planning in Task 2 and for building the CI/CD pipeline in the later tasks."),
    ],
)

TASKS[2] = dict(
    title="Agile Planning and DevOps Workflow",
    branch="task/02-agile-planning",
    filename="task-02-agile-planning",
    blocks=[
        ("h", "1. Objective"),
        ("p", "To create user stories and acceptance criteria for the Parcel Tracking Portal, prepare a product backlog, define a 15-task Kanban plan and a Definition of Done, and show the DevOps lifecycle from development to operations."),
        ("h", "2. Agile Approach"),
        ("p", "Kanban is used because the project is a fixed sequence of 15 tasks delivered by one person. Work is visualised on a GitHub Projects board with the columns Backlog, To Do, In Progress, In Review and Done. A work-in-progress (WIP) limit of one task In Progress at a time keeps the flow steady."),
        ("h", "3. User Stories and Acceptance Criteria"),
        ("table", ["ID", "User story", "Acceptance criteria"], [
            ["US-01", "As a warehouse staff member, I want to add a new parcel so that it can be tracked.", "Form accepts parcel ID, sender, receiver, origin and destination. Duplicate parcel IDs are rejected. A success message is shown after saving."],
            ["US-02", "As a delivery agent, I want to add a status event to a parcel so that its progress is recorded.", "Status can be Picked Up, In Transit, Out for Delivery, Delivered or Failed. Each event stores a timestamp and location. The latest status updates on the dashboard."],
            ["US-03", "As a customer, I want to search a parcel by ID so that I can see its current status.", "Search returns the result in under 2 seconds. An unknown ID shows a clear 'not found' message."],
            ["US-04", "As a support agent, I want a dashboard listing all parcels so that I can monitor them.", "Dashboard lists all parcels with ID, status and last update. The list can be filtered by status."],
            ["US-05", "As a manager, I want summary indicators so that I can see overall performance.", "Counts of total, in transit, delivered and delayed parcels are shown and match the stored data."],
            ["US-06", "As a customer, I want to open a parcel and see its full history so that I know what happened.", "Drill-down page shows all events in time order on one screen."],
            ["US-07", "As a support agent, I want delayed or failed parcels flagged so that I can act quickly.", "A parcel with no update beyond the set limit, or with status Failed, appears in the alert view."],
            ["US-08", "As a developer, I want every commit built and tested automatically so that defects are found early.", "A Jenkins build is triggered on commit. A failed test stops deployment. Reports are archived."],
            ["US-09", "As an operations engineer, I want the portal deployed in a container and provisioned by script so that environments are repeatable.", "The Docker image runs the portal. The Ansible/Puppet script builds the environment and a second run changes nothing."],
        ]),
        ("h", "4. Product Backlog"),
        ("table", ["ID", "Backlog item", "Priority", "Effort", "Linked task"], [
            ["B-01", "Problem definition and scope", "High", "S", "Task 1"],
            ["B-02", "Agile plan, backlog and workflow", "High", "S", "Task 2"],
            ["B-03", "Architecture and technology setup", "High", "M", "Task 3"],
            ["B-04", "Repository initialization and conventions", "High", "S", "Task 4"],
            ["B-05", "Parcel entry and event entry (US-01, US-02)", "High", "M", "Task 5"],
            ["B-06", "Search, dashboard, summary, drill-down, alerts (US-03 to US-07)", "High", "L", "Task 6"],
            ["B-07", "Jenkins installation and CI job (US-08)", "High", "M", "Task 7"],
            ["B-08", "Jenkinsfile and deployment to Tomcat", "High", "M", "Task 8"],
            ["B-09", "Selenium test cases for three user journeys", "High", "M", "Task 9"],
            ["B-10", "Selenium tests integrated in Jenkins", "Medium", "M", "Task 10"],
            ["B-11", "Dockerfile and container lifecycle (US-09)", "Medium", "M", "Task 11"],
            ["B-12", "Jenkins to Docker continuous deployment", "Medium", "L", "Task 12"],
            ["B-13", "Ansible/Puppet configuration script", "Medium", "M", "Task 13"],
            ["B-14", "Automated provisioning and health check", "Medium", "M", "Task 14"],
            ["B-15", "Final release, documentation and viva", "High", "M", "Task 15"],
        ]),
        ("p", "Effort key: S = small, M = medium, L = large."),
        ("h", "5. 15-Task Kanban Plan"),
        ("table", ["Task", "Title", "Phase", "Board column (at start)"], [
            ["1", "Problem Definition and Scope", "Plan", "Done"],
            ["2", "Agile Planning and DevOps Workflow", "Plan", "In Progress"],
            ["3", "Requirements, Architecture and Technology Setup", "Plan", "To Do"],
            ["4", "Git and GitHub Repository Initialization", "Code", "To Do"],
            ["5", "Feature Development with Branching", "Code", "Backlog"],
            ["6", "MVP Completion and Git Collaboration", "Code", "Backlog"],
            ["7", "Jenkins Installation and CI Job", "Build", "Backlog"],
            ["8", "Pipeline as Code and Server Deployment", "Build / Deploy", "Backlog"],
            ["9", "Selenium Test Design and Local Execution", "Test", "Backlog"],
            ["10", "Continuous Testing in Jenkins", "Test", "Backlog"],
            ["11", "Docker Image and Container Lifecycle", "Release", "Backlog"],
            ["12", "Jenkins-Docker Continuous Deployment", "Deploy", "Backlog"],
            ["13", "Configuration Management Script", "Operate", "Backlog"],
            ["14", "Automated Provisioning and Reliability Validation", "Operate / Monitor", "Backlog"],
            ["15", "Final End-to-End Release, Documentation and Viva", "Release / Monitor", "Backlog"],
        ]),
        ("h2", "Kanban rules"),
        ("bullets", [
            "Only one task is In Progress at a time (WIP limit = 1).",
            "A task moves to In Review when its pull request is opened.",
            "A task moves to Done only when it meets the Definition of Done.",
            "Each task has its own branch and its own report.",
        ]),
        ("h", "6. Definition of Done"),
        ("p", "A task is Done only when all of the following are true:"),
        ("check", [
            "Work is committed on its own task branch with meaningful commit messages.",
            "Code or configuration builds and runs without errors.",
            "Acceptance criteria of the related user stories are met.",
            "Tests (where applicable) pass.",
            "Pull request is raised, reviewed and merged into develop.",
            "The branch is kept for evidence and not deleted.",
            "Task report (Word) is written with commands, output and screenshots.",
            "The Kanban card is moved to Done.",
        ]),
        ("h", "7. DevOps Lifecycle: Development to Operations"),
        ("p", "The Parcel Tracking Portal follows a continuous loop. Code changes flow from planning through build, test, release and deployment into operations, and monitoring feedback returns to planning."),
        ("table", ["Stage", "Activity in this project", "Tool"], [
            ["1. Plan", "User stories, backlog, Kanban board", "GitHub Projects, Issues"],
            ["2. Code", "Feature branches, commits, pull requests", "Git, GitHub, VS Code"],
            ["3. Build", "Compile and package the WAR file", "Maven, Jenkins"],
            ["4. Test", "Automated UI tests on every build", "Selenium, JUnit, Jenkins"],
            ["5. Release", "Versioned Docker image", "Docker, Docker Hub"],
            ["6. Deploy", "Automatic deployment to Tomcat or container", "Jenkins, Docker"],
            ["7. Operate", "Provision and configure servers", "Ansible / Puppet"],
            ["8. Monitor", "Health check, test reports, feedback to Plan", "Jenkins reports, curl health check"],
        ]),
        ("code", "PLAN -> CODE -> BUILD -> TEST -> RELEASE -> DEPLOY -> OPERATE -> MONITOR\n  ^                                                                |\n  +------------------------- feedback loop ------------------------+"),
        ("h", "8. Steps Performed"),
        ("numbered", [
            "Created branch task/02-agile-planning from develop.",
            "Wrote user stories with acceptance criteria.",
            "Prepared the product backlog and mapped it to the 15 tasks.",
            "Created the Kanban board on GitHub Projects with five columns.",
            "Defined the Definition of Done and the DevOps lifecycle.",
            "Committed the report, raised a pull request into develop and merged it.",
        ]),
        ("h", "9. Git Commands Used"),
        ("code", "git checkout develop\ngit pull origin develop\ngit checkout -b task/02-agile-planning\ngit add docs\ngit commit -m \"docs: add Task 2 agile planning report\"\ngit push -u origin task/02-agile-planning"),
        ("h", "10. Screenshots"),
        ("screenshots",),
        ("h", "11. Deliverables Checklist"),
        ("check", [
            "User stories with acceptance criteria",
            "Product backlog",
            "15-task Kanban plan",
            "Definition of Done",
            "DevOps lifecycle diagram",
        ]),
        ("h", "12. Issues Faced and Fixes"),
        ("bullets", [
            "Script initially run from the wrong folder; fixed by running it inside the repository.",
            "Report file was open in Word and caused a permission error; closed it before regenerating.",
        ]),
        ("h", "13. Conclusion"),
        ("p", "Agile planning for the Parcel Tracking Portal is complete. The user stories, backlog, Kanban board, Definition of Done and DevOps lifecycle give a clear path for building the application and its CI/CD pipeline in the next tasks."),
    ],
)

# ---------- RUN ----------
if __name__ == "__main__":
    nums = [int(a) for a in sys.argv[1:]] or sorted(TASKS)
    for n in nums:
        if n not in TASKS:
            print(f"Task {n} content not added yet."); continue
        t = TASKS[n]
        build(n, t["title"], t["branch"], t["blocks"], t["filename"])